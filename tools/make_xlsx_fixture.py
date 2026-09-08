#!/usr/bin/env python3
"""시험용 xlsx 고정 자료를 다시 만든다.

  python3 tools/make_xlsx_fixture.py [원본.xlsx]

tests/test_attach.mjs가 쓰는 tests/fixtures/real_excel.xlsx를 만든다. 손으로 조립한
최소 OOXML 자료와 달리, 이 파일은 엑셀 라이브러리(openpyxl)가 써 낸 것이라 시트
이름·열 배치·빈 칸이 실제 문서 그대로다. 표가 어긋나는지 보는 데 쓴다.

원본 「지역사회보장지표.xlsx」는 192시트 7.6MB(대부분 styles.xml)라 저장소에 두지
않는다. 앞 세 시트의 40행만 옮겨 16KB로 줄인다. 값은 이미 app/data/catalog.json으로
공개돼 있는 지표 메타다.
"""

from __future__ import annotations

import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
DEFAULT_SRC = Path("/home/user/kihasa-indicator-new/source/지역사회보장지표.xlsx")
DST = ROOT / "tests" / "fixtures" / "real_excel.xlsx"

SHEETS = 3        # 앞에서 몇 장을 옮길지
ROWS = 40         # 시트마다 몇 행까지
CELL_CHARS = 200  # 셀 하나가 너무 길면 자른다


def main(argv: list[str]) -> int:
    try:
        import openpyxl
    except ModuleNotFoundError:
        print("[중단] openpyxl이 필요하다 — pip install openpyxl", file=sys.stderr)
        return 1

    src = Path(argv[0]) if argv else DEFAULT_SRC
    if not src.exists():
        print(f"[중단] 원본이 없다: {src}", file=sys.stderr)
        return 1

    wb = openpyxl.load_workbook(src, read_only=True, data_only=True)
    out = openpyxl.Workbook()
    out.remove(out.active)
    for name in wb.sheetnames[:SHEETS]:
        ws = out.create_sheet(name[:31])
        for r, row in enumerate(wb[name].iter_rows(max_row=ROWS, values_only=True), start=1):
            for c, v in enumerate(row, start=1):
                if v is None:
                    continue
                ws.cell(row=r, column=c,
                        value=v[:CELL_CHARS] if isinstance(v, str) else v)
        print(f"  {name}: {ws.max_row}행 × {ws.max_column}열")

    DST.parent.mkdir(parents=True, exist_ok=True)
    out.save(DST)
    print(f"[됨] {DST.relative_to(ROOT)} — {DST.stat().st_size:,}바이트")
    return 0


if __name__ == "__main__":
    raise SystemExit(main(sys.argv[1:]))
